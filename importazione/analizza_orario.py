
from collections import Counter
from openpyxl import load_workbook
import re

GIORNI = [
    ("Lunedì", 1, 7),
    ("Martedì", 8, 7),
    ("Mercoledì", 15, 6),
    ("Giovedì", 21, 7),
    ("Venerdì", 28, 7),
]


def normalizza_classe(valore):
    testo = str(valore).strip()
    testo = re.sub(r"\s*\n\s*", "", testo)

    correzioni = {
        "1BA": "1BAF",
        "2BA": "2BAF",
    }

    return correzioni.get(testo, testo)


def analizza_file(file_excel):
    workbook = load_workbook(file_excel, data_only=True)
    sheet = workbook.active

    lezioni = []
    docenti = set()
    valori_speciali = Counter()

    try:
        for numero_riga in range(4, sheet.max_row + 1):
            docente = sheet.cell(numero_riga, 1).value

            if docente is None:
                continue

            docente = str(docente).strip()

            if not docente or docente.casefold() == "docente":
                continue

            docenti.add(docente)

            for giorno, colonna_iniziale, numero_ore in GIORNI:
                for ora in range(1, numero_ore + 1):
                    colonna = colonna_iniziale + ora
                    valore = sheet.cell(numero_riga, colonna).value

                    if valore is None or not str(valore).strip():
                        continue

                    classe = normalizza_classe(valore)

                    lezioni.append({
                        "docente": docente,
                        "giorno": giorno,
                        "ora": ora,
                        "classe": classe,
                        "riga_excel": numero_riga,
                    })

                    if classe.upper() in {"D", "(D)"}:
                        valori_speciali[classe.upper()] += 1
    finally:
        workbook.close()

    return {
        "lezioni": lezioni,
        "numero_docenti": len(docenti),
        "numero_lezioni": len(lezioni),
        "valori_speciali": dict(valori_speciali),
    }

def crea_righe_anteprima(lezioni):
    righe = {}

    for lezione in lezioni:
        chiave = (lezione["docente"], lezione["giorno"])

        if chiave not in righe:
            righe[chiave] = {
                "docente": lezione["docente"],
                "giorno": lezione["giorno"],
                "ore": [""] * 7,
            }

        indice = lezione["ora"] - 1

        if 0 <= indice < 7:
            righe[chiave]["ore"][indice] = lezione["classe"]

    
    ordine_giorni = {
        "Lunedì": 1,
        "Martedì": 2,
        "Mercoledì": 3,
        "Giovedì": 4,
        "Venerdì": 5,
    }

    return sorted(
        righe.values(),
        key=lambda riga: (
            riga["docente"].casefold(),
            ordine_giorni.get(riga["giorno"], 99),
        ),
    )
    

def controlla_conflitti(lezioni):
    assegnazioni = {}
    conflitti = []

    for lezione in lezioni:
        chiave = (
            lezione["docente"],
            lezione["giorno"],
            lezione["ora"],
        )
        classe = lezione["classe"]

        if chiave in assegnazioni:
            classe_precedente = assegnazioni[chiave]

            if classe != classe_precedente:
                conflitti.append({
                    "docente": lezione["docente"],
                    "giorno": lezione["giorno"],
                    "ora": lezione["ora"],
                    "classe1": classe_precedente,
                    "classe2": classe,
                    "riga_excel": lezione["riga_excel"],
                })
        else:
            assegnazioni[chiave] = classe

    return conflitti
    

def importa_lezioni(lezioni):
    from collections import defaultdict
    from django.db import connection, transaction

    conflitti = controlla_conflitti(lezioni)

    if conflitti:
        raise ValueError(
            f"Importazione annullata: trovati {len(conflitti)} conflitti."
        )

    orari = defaultdict(lambda: {str(ora): "" for ora in range(1, 8)})

    for lezione in lezioni:
        nome = lezione["docente"]
        giorno = lezione["giorno"]
        ora = str(lezione["ora"])
        classe = lezione["classe"]

        orari[(nome, giorno)][ora] = classe

    with transaction.atomic():
        with connection.cursor() as cursor:
            cursor.execute("DELETE FROM public.orario")

            for (nome, giorno), ore in orari.items():
                cursor.execute("""
                    INSERT INTO public.orario
                    (nome, giorno, "1", "2", "3", "4", "5", "6", "7")
                    VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
                """, [
                    nome,
                    giorno,
                    ore["1"],
                    ore["2"],
                    ore["3"],
                    ore["4"],
                    ore["5"],
                    ore["6"],
                    ore["7"],
                ])

    return len(orari)



