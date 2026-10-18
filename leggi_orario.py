
from collections import Counter
from openpyxl import load_workbook
import re

FILE_EXCEL = "dati/orario2024.xlsx"

GIORNI = [
    ("Lunedì", 1, 7),
    ("Martedì", 8, 7),
    ("Mercoledì", 15, 6),
    ("Giovedì", 21, 7),
    ("Venerdì", 28, 7),
]


def normalizza_classe(valore):
    """Normalizza le sigle delle classi secondo regole esplicite."""
    testo = str(valore).strip()
    testo = re.sub(r"\s*\n\s*", "", testo)

    correzioni = {
        "1BA": "1BAF",
        "2BA": "2BAF",
    }

    return correzioni.get(testo, testo)


workbook = load_workbook(FILE_EXCEL, data_only=True)
sheet = workbook.active

lezioni = []
docenti = set()
valori_speciali = Counter()

# Le prime tre righe contengono intestazioni e numeri delle ore.
for numero_riga in range(4, sheet.max_row + 1):
    docente = sheet.cell(numero_riga, 1).value

    if docente is None:
        continue

    docente = str(docente).strip()

    # Ignora le intestazioni ripetute nel foglio.
    if not docente or docente.casefold() == "docente":
        continue

    docenti.add(docente)

    for giorno, colonna_iniziale, numero_ore in GIORNI:
        for ora in range(1, numero_ore + 1):
            colonna = colonna_iniziale + ora
            valore = sheet.cell(numero_riga, colonna).value

            if valore is None or not str(valore).strip():
                continue

            classe = normalizza_classe(valore)

            lezione = {
                "docente": docente,
                "giorno": giorno,
                "ora": ora,
                "classe": classe,
                "riga_excel": numero_riga,
            }

            lezioni.append(lezione)

            if classe.upper() in {"D", "(D)"}:
                valori_speciali[classe.upper()] += 1

workbook.close()

print(f"Docenti trovati: {len(docenti)}")
print(f"Ore compilate: {len(lezioni)}")
print(f"Valori D e (D): {dict(valori_speciali)}")

print("\nPrime 30 assegnazioni:")
for lezione in lezioni[:30]:
    print(
        f"{lezione['docente']:25} "
        f"{lezione['giorno']:12} "
        f"ora {lezione['ora']} "
        f"classe={lezione['classe']}"
    )

print("\nNumero di assegnazioni per giorno:")
conteggio = Counter(x["giorno"] for x in lezioni)

for giorno, _, _ in GIORNI:
    print(f"{giorno}: {conteggio[giorno]}")

print("\nControllo completato: nessun dato è stato scritto nel database.")

print("\nValori delle classi presenti nel file:")
classi_trovate = sorted({x["classe"] for x in lezioni})

for classe in classi_trovate:
    print(classe)
    
import os
import django

os.environ.setdefault("DJANGO_SETTINGS_MODULE", "config.settings")
django.setup()

from django.db import connection

with connection.cursor() as cursor:
    cursor.execute("SELECT classe FROM classi")
    classi_database = {riga[0].strip() for riga in cursor.fetchall()}

classi_excel = {x["classe"].strip() for x in lezioni}

valori_speciali = {"D", "(D)", "R"}

print("\n--- CLASSI PRESENTI NELL'EXCEL MA NON NEL DATABASE ---")
for classe in sorted(classi_excel - classi_database - valori_speciali):
    print(classe)

print("\n--- VALORI SPECIALI NELL'EXCEL ---")
for classe in sorted(classi_excel & valori_speciali):
    print(classe)

print("\n--- CLASSI DELL'EXCEL GIÀ PRESENTI NEL DATABASE ---")
for classe in sorted(classi_excel & classi_database):
    print(classe)

